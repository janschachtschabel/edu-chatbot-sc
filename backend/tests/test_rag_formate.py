"""Paket 7, P6 (P5-8): PDF, Word, PowerPoint und Excel in die Wissensbasis.

Bis Paket 7 lag markitdown ohne Extras im Image; diese vier Formate scheiterten
mit 415 (Paket 4, C5). Hier läuft der ECHTE Wandler, nichts ist gefälscht: der
Test fällt, wenn ein Extra im Lock fehlt — das Image baut aus demselben Lock
(``uv sync --locked --no-dev``).

Die Dateien entstehen im Test, klein und ohne Netz. DOCX und PDF von Hand (kein
Schreib-Paket nötig), PPTX und XLSX mit den Paketen, die das jeweilige Extra
ohnehin mitbringt.
"""

from __future__ import annotations

import asyncio
import io
import zipfile
from pathlib import Path

from boerdi.services.rag.umwandeln import Umwandlungsfehler, convert_to_markdown

SATZ = "Fehlerkultur im Unterricht"


def _docx(pfad: Path) -> None:
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    teile = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType='
            '"application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>"),
        "_rels/.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>"),
        "word/document.xml": (
            f'<w:document xmlns:w="{w}"><w:body><w:p><w:r><w:t>{SATZ}</w:t></w:r></w:p>'
            "</w:body></w:document>"),
    }
    with zipfile.ZipFile(pfad, "w") as z:
        for name, inhalt in teile.items():
            z.writestr(name, inhalt)


def _pdf(pfad: Path) -> None:
    """Ein einseitiges PDF mit einer Textzeile — Querverweise mit echten Abständen."""
    strom = f"BT /F1 18 Tf 20 100 Td ({SATZ}) Tj ET".encode("latin-1")
    objekte = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(strom) + strom + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    abstaende = []
    for nr, obj in enumerate(objekte, start=1):
        abstaende.append(out.tell())
        out.write(b"%d 0 obj\n" % nr + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objekte) + 1))
    for a in abstaende:
        out.write(b"%010d 00000 n \n" % a)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
              % (len(objekte) + 1, xref))
    pfad.write_bytes(out.getvalue())


def _pptx(pfad: Path) -> None:
    from pptx import Presentation  # kommt mit dem Extra ``pptx``

    praesentation = Presentation()
    folie = praesentation.slides.add_slide(praesentation.slide_layouts[1])
    folie.shapes.title.text = SATZ
    praesentation.save(pfad)


def _xlsx(pfad: Path) -> None:
    from openpyxl import Workbook  # kommt mit dem Extra ``xlsx``

    mappe = Workbook()
    blatt = mappe.active
    blatt.append(["Thema", "Stufe"])
    blatt.append([SATZ, "Sek I"])
    mappe.save(pfad)


def _wandeln(pfad: Path) -> str:
    text = asyncio.run(convert_to_markdown(str(pfad)))
    assert not isinstance(text, Umwandlungsfehler), text
    return text


def test_word_wird_gelesen(tmp_path) -> None:
    datei = tmp_path / "unterricht.docx"
    _docx(datei)
    assert SATZ in _wandeln(datei)


def test_pdf_wird_gelesen(tmp_path) -> None:
    datei = tmp_path / "unterricht.pdf"
    _pdf(datei)
    assert SATZ in _wandeln(datei)


def test_powerpoint_wird_gelesen(tmp_path) -> None:
    datei = tmp_path / "unterricht.pptx"
    _pptx(datei)
    assert SATZ in _wandeln(datei)


def test_excel_wird_gelesen_mit_tabelle(tmp_path) -> None:
    datei = tmp_path / "unterricht.xlsx"
    _xlsx(datei)
    text = _wandeln(datei)
    assert SATZ in text
    assert "Sek I" in text
    assert "|" in text  # als Markdown-Tabelle, nicht als Fließtext


def test_ein_link_auf_ein_pdf_wird_gelesen(tmp_path, monkeypatch) -> None:
    """„Webseite einlesen“ mit einem Link auf ein PDF: bis P6 kam 415 zurück.
    Der Abruf ist gefälscht (offen, ohne Netz), der Wandler echt; die numerische
    Adresse braucht kein DNS — wie in ``test_rag_umwandeln``."""
    from boerdi.services import url_safety
    from boerdi.services.rag.umwandeln import convert_url_to_markdown

    datei = tmp_path / "heft.pdf"
    _pdf(datei)

    async def _abruf(url, *, max_bytes, frist_s):
        return datei.read_bytes(), "application/pdf"

    monkeypatch.setattr(url_safety, "hole_begrenzt", _abruf)
    text = asyncio.run(convert_url_to_markdown("http://8.8.8.8/heft.pdf"))
    assert not isinstance(text, Umwandlungsfehler), text
    assert SATZ in text
